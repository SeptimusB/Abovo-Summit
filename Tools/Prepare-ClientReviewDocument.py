"""Create a numbered client review without overwriting the source or an earlier copy."""
import argparse
from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

ORANGE = "E97132"
FIXES = {
    7: "The non-working Start Date Edit link has been removed.",
    8: "The lower Funding Status figures have been realigned.",
    12: "Save and Save As are available on the DIT toolbar. Save is disabled when the file is clean.",
    13: "Navigator captions have been shortened and width behaviour revised. Please confirm this at your Windows display scale.",
    16: "Enter commits and advances on the last horizontal or vertical navigation axis; Shift+Enter reverses it. Traversal stays within the current tab, including adjacent grids and header inputs.",
    17: "A single copied dropdown value can fill selected cells. Invalid dropdown values reject the whole paste; grouped Undo restores the previous values.",
    18: "Check Sheet is under Global Assumptions. Normal Save and Save As remain available with a warning; recovery autosaves have a separate configurable Check Sheet policy.",
    29: "Repeated edits no longer refit and enlarge the edited column. Please retest the Voids table.",
    34: "Monetary display now follows the source workbook cell, including the two-decimal weekly service charge.",
    42: "The shared post-edit column resizing has been removed. Please retest this table.",
    48: "Repeated dropdown and numeric edits preserve the chosen column width.",
    53: "Empty survey dates remain blank instead of displaying 30/12/1899. Date entry, clearing and Undo/Redo passed native tests.",
    62: "The second caption now reads Period Units into Mgmt (to), matching the existing workbook input.",
    67: "SHG Profiling and SHG Calculation Basis now use their separate workbook inputs; edits and Undo were tested independently.",
    69: "Weekly rent and service charge display use the source cell's two-decimal format.",
    77: "Monetary Development Expenditure inputs retain decimal values through entry, paste and save/reopen. Year headers remain integers.",
    84: "Percentage editors no longer impose an invented zero-to-100-percent limit. Negative values are allowed where the workbook and explicit field rules permit them.",
    95: "Funding input refresh and keyboard navigation preserve focus and scroll position in native tests. Please confirm with the client's physical keyboard.",
    101: "First Interest Payment Month writes an actual month-end date, displayed as a month, rather than month-name text. Validation, blank values and Undo/Redo were tested.",
    108: "BP Year and Year now read the correct workbook columns.",
    109: "Covenant input columns are editable according to source protection; the two year columns remain read-only. Monetary values retain decimals.",
    120: "Monetary Housing Asset inputs retain decimals. Counts, years and useful-life periods intentionally retain their existing types.",
    127: "The amount input retains decimals through native editing, Undo/Redo and XLSB save/reopen.",
    130: "Journal Amount retains decimals; Year remains an integer. Editing and save/reopen were tested.",
    132: "Journals insertion updates IR_Journals and both Transactional DB mirrors. Add/delete, save/reopen and damaged-name rejection tests passed on disposable files.",
}
NOTES = {
    4: "Question: do Open, New and Compare now fit their circles on the client's display? Please give Windows scaling and a screenshot if not. Existing orange is retained; client DPI acceptance is still needed.",
    5: "Question: which icon still overflows, if any, and at what Windows scaling? The original orange status is retained, not a new physical-display test result.",
    6: "Question: is Program Information fully visible in both restored and maximised windows now? The original orange status is retained pending that check.",
    9: "Question: which summary container still has an unnecessary horizontal scrollbar? Please include a restored-window screenshot. Original orange is retained; this item is not independently closed.",
    10: "Question: after changing Company Name in Global Assumptions, which heading fails to update, and does reopening that interface update it?",
    11: "Question: which button is blank in the current version? Please identify its neighbours or attach a screenshot; the toolbar has changed since 2.51.",
    14: "Question: should every repeating year header display Year 1, Year 2, etc., rather than bare numbers? Please distinguish the displayed label from the numeric stored year.",
    15: "Question: which current button is intended to return to the preceding interface? Please identify it so it is not confused with Undo, History or Rebuild.",
    21: "Question: does Company Name fail to accept typing, lose focus, or fail to refresh the heading after Enter/Tab? Please give the exact sequence.",
    22: "Review: remembered Check Sheet warnings, selected-model integrity checks and unsaved-session warning retention have since been repaired. This does not establish that every numerical discrepancy is resolved. Question: if it persists, which workbook, Check Sheet row and Summit/Excel values disagree after a fresh check of that same file?",
    25: "Question: what description must fit without truncation? Please provide a screenshot at the client's normal window size and display scale.",
    26: "Question: should the gap be removed entirely, or retained as a small separation between those two tables?",
    39: "Question: should Summit omit Add Lines for this category table, or is a new supported expansion action wanted? The absence of a master macro is not assumed to authorise a new structural rule.",
    45: "Question: which default category and expected colour/font differ from Excel? Please give a before/after example.",
    46: "Open: the workbook conditional rule can change after Description is cleared while the DIT colour remains stale. This refresh repair is still outstanding.",
    47: "Question: which workbook, row/category and preceding action should lock Year 1? The sampled source cells are unlocked; a blanket interface lock has not been applied.",
    49: "Open: this means pinning Description during horizontal scrolling, not preventing edits. Question: should only Description be pinned, or the category columns too?",
    52: "Open: conditional-format refresh, truncated descriptions and the delayed three-table refresh remain separate layout/refresh work. Question: which three tables require switching away, and what input triggers it?",
    54: "Partly repaired: Survey cost inputs now retain decimal entry and paste values. Header wrapping, widths and nested scrolling are still open. Question: please retest a rectangular paste from Excel and identify any failing destination cells.",
    57: "Question: is the missing scroll vertical or horizontal, and which Import and Edit Schemes control is inaccessible?",
    58: "Question: please supply the exact scheme/workbook and whether Scheme Name is blank, locked, or overwritten after entry. The Unidentified default needs a precise reproduction.",
    60: "Question: is the request only to rename the two captions, or is the Scheme Name value also missing? Please give a scheme example for the latter.",
    63: "Open: workbook-based conditional-format refresh is still to be repaired. Question: please give the intended appearance before and after House Type is entered.",
    64: "Open: Include-dependent formatting is not marked fixed. Question: should excluded fields be grey but editable, or locked as well, matching which source workbook rule?",
    68: "Question: what exact display is required for Sales Value per Unit: currency symbol, grouping and how many decimals? Please show the matching Excel cell.",
    70: "Question: are the Void Rate and Bad Debt Rate From Year headings wrong, the editors missing, or the stored values incorrect?",
    78: "Open: the Planned Maintenance repeating year-header editor definition is incomplete; the monetary body-input fix does not repair this header.",
    79: "Question: should the label From Year be used consistently, including the other tables mentioned under General?",
    83: "Open: conditional formatting and nested scrolling are not covered by the negative-percentage input fix.",
    85: "Question: which Economic Assumptions field and change should alter its appearance? Please provide the starting value, new value and expected colour or screenshot.",
    87: "Question: which workbook/version shows Service Costs here? The current definition points to the Development Repairs sources; please provide a screenshot of the mismatch.",
    90: "Question: is the missing highlight on a selected cell before editing, on the open editor, or on selected pasted cells?",
    93: "Question: should the horizontal scrollbar stay visible at the bottom of the window while vertically scrolling the Funding section?",
    94: "Question: should Loan Description be frozen as a header across all Funding tables while scrolling vertically? Please identify any other rows that must remain visible.",
    97: "Question: which entries are offered now, and what entries should Facility Name offer? Please supply the corresponding Excel list or a specific facility example.",
    99: "Question: does date ghost mean a format hint in an empty input, a blank instead of a zero date, or a calculated default shown in grey?",
    100: "Question: for Fixed until and including, which of those date behaviours is intended? Please show an empty and a populated example.",
    104: "Question: which exact Variable or Cash Rates table and first date should be locked? A global first-date lock has not been assumed.",
    105: "Open: fee-description editors and dependent eligibility need work. AGL also has shifted Other Fees coordinates; the current-template BP-year correction does not repair that older layout. Question: please identify the workbook and fee table to use for client acceptance. No forced unlocking has been applied.",
    116: "Question: should years run down rows and cost categories across columns after transposing? Please confirm the desired orientation with an example.",
    117: "Open: transpose/layout and the Repairs caption remain unverified. Question: should the orientation match the answer for Management Costs Capitalisation above?",
    121: "Question: are component descriptions already defined in this workbook, or should the interface create additional component types? Please give a concrete component and intended source location.",
    122: "Question: should Grant Assumptions contain only grant inputs, with depreciation and useful-life inputs together under Housing Asset Depreciation? Current definitions overlap, so labels have not simply been swapped.",
    123: "Question: should Remaining Useful Life remain a separate tab after the grouping above, or be included within Housing Asset Depreciation?",
}


def note_after(source, text, orange=False):
    node = OxmlElement("w:p")
    source._p.addnext(node)
    p = Paragraph(node, source._parent)
    p.style = "Normal"
    p.paragraph_format.left_indent = Pt(14)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(ORANGE if orange else "333333")
    return p


def colour_tail(p, marker):
    """Split only intersecting runs, retaining every source character and run style."""
    start = p.text.index(marker)
    position = 0
    for run in list(p.runs):
        end = position + len(run.text)
        if end > start:
            if position < start:
                tail = deepcopy(run._r)
                run._r.addnext(tail)
                from docx.text.run import Run
                tail_run = Run(tail, p)
                tail_run.text = run.text[start - position:]
                run.text = run.text[:start - position]
                tail_run.font.color.rgb = RGBColor.from_string(ORANGE)
            else:
                run.font.color.rgb = RGBColor.from_string(ORANGE)
        position = end


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--revision", choices=("01", "02"), default="02")
    args = parser.parse_args()
    release = "2.82" if args.revision == "01" else "2.83"
    retest = {}
    if args.revision == "02":
        FIXES.update({
            10: "Changing Setup Company Name now refreshes the current company headings, main file tab and FileInstance summary after committing the edit. Undo and Redo refresh them too. The name comes from the workbook's SelectTrust input.",
            14: "Repeating ordinal-year headers consistently show Year 1, Year 2, etc. Dropdowns retain the financial-year suffix where available. Only the presentation changes; the stored year remains numeric.",
            15: "The linked-interface Return button now uses a left arrow and appears at the far right after its own separator. History uses the former curved Return icon. Return and its separator disappear together when there is no return link.",
        })
        NOTES.update({
            4: "> Alex to test.",
            5: "> Alex to confirm on his PC.",
            9: "> Alex to confirm. HTML reformatting is believed to have removed the unnecessary horizontal scrollbars.",
        })
        for number in (10, 14, 15):
            NOTES.pop(number, None)
        retest = {number: "> Alex to confirm." for number in (7, 8, 10, 12, 13, 14, 15)}
    if args.output.exists() or args.source.resolve() == args.output.resolve():
        raise SystemExit("Refusing to overwrite the source or an earlier revision")
    source_hash = sha256(args.source.read_bytes()).hexdigest()
    doc = Document(args.source)
    originals = list(doc.paragraphs)
    original_text = [p.text for p in originals]
    if len(originals) != 137 or "IR_Journals" not in originals[131].text:
        raise SystemExit("Source revision differs; recheck paragraph mappings before authoring")

    title = originals[0].insert_paragraph_before("Summit client issue review version " + args.revision, "Title")
    for r in title.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)
    intro = originals[0].insert_paragraph_before(
        "Review against Summit " + release + " on 23 September 2026. Please retest the orange items and answer the questions beneath the relevant comments. The original 2.51 wording and order are retained."
        + (" Alex is the Abovo director responsible for the client confirmations below. Alex to confirm means acceptance is pending, not already given." if args.revision == "02" else "")
    )
    intro.paragraph_format.space_after = Pt(8)
    legend = originals[0].insert_paragraph_before(
        "Orange means believed fixed, not client or accountant acceptance. Earlier orange markings are retained, with qualifications where confirmation is still needed. Green retains its original meaning of agreed, not necessarily fixed. Uncoloured issues remain open or unverified; a repaired part does not close the rest of a mixed item."
    )
    legend.paragraph_format.space_after = Pt(12)
    for number, p in enumerate(originals, 1):
        if p.text.strip() and all(r.bold for r in p.runs if r.text.strip()):
            p.paragraph_format.keep_with_next = True
        if number in FIXES:
            for r in p.runs:
                r.font.color.rgb = RGBColor.from_string(ORANGE)
            p.paragraph_format.keep_with_next = True
            note_after(p, "Believed fixed: " + FIXES[number] + " " + retest.get(number, "Please confirm in the client retest."))
        elif number in NOTES:
            p.paragraph_format.keep_with_next = True
            note_after(p, NOTES[number])
    colour_tail(originals[53], "Can’t seem to paste")
    assert [p.text for p in originals] == original_text, "A source comment changed"
    # The source ends with five empty paragraphs, including a manual page break.
    # They carry no comments and would create a blank final page in this revision.
    for p in reversed(originals):
        if p.text.strip():
            break
        p._p.getparent().remove(p._p)
    doc.core_properties.title = "Summit client issue review version " + args.revision
    doc.core_properties.subject = "Client retest and clarification against Summit " + release
    doc.core_properties.comments = "Numbered review copy. Original client report preserved."
    doc.save(args.output)
    assert sha256(args.source.read_bytes()).hexdigest() == source_hash
    check = Document(args.output)
    assert all(text in [p.text for p in check.paragraphs] for text in original_text)
    print(f"Saved {args.output}; {len(FIXES)} reviewed fixed items, one partly repaired item, {len(NOTES)} review notes; original SHA256 {source_hash}")


if __name__ == "__main__":
    main()
