"""Produce review v05; retain earlier evidence and distinguish Jon/Alex acceptance."""
from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt
from docx.shared import RGBColor
from docx.oxml.ns import qn
import sys
import importlib.util

spec = importlib.util.spec_from_file_location('stages', Path(__file__).with_name('Update-ClientReviewStages.py'))
stages = importlib.util.module_from_spec(spec)
spec.loader.exec_module(stages)
source = Path('Library/Client_Review_2.85_v04.docx')
output = Path('Library/Client_Review_2.93_v05.docx')
if output.exists() and '--replace-draft' not in sys.argv:
    raise SystemExit('Choose a new revision; output already exists.')
before = source.read_bytes()
d = Document(source)
p = list(d.paragraphs)
assert len(p) == 205 and 'version 04' in p[0].text

def setp(i, text, fill=None, colour=stages.DARK):
    stages.replace(p[i], text)
    stages.format_text(p[i], colour, fill)

def jon(i, text):
    setp(i, text + '\n> Jon to test', stages.BLUE)

def alex(heading, detail, text):
    stages.format_text(p[heading], stages.ORANGE)
    setp(detail, text + '\n> Alex to test — Test response 2', stages.LIGHT_BLUE)

setp(0, 'Summit client issue review version 05')
setp(1, 'Ready to test with Summit 2.93\n24 September 2026\n\nThis revision records Jon’s latest results and the simpler Funding scheduler requested with Alex. Save/break and file-warning tests have passed Jon’s functional test; the live Check Sheet header and Transactional DB concern remain open. Earlier issue order and screenshots are retained.')
setp(3, 'Blue: Jon to test — new or previously unconfirmed changes still need Jon’s functional test.', stages.BLUE)
setp(4, 'Orange issue heading + light-blue detail: Alex to test. Test response 2 identifies the items Jon has just passed to Alex; it is not Alex’s sign-off.', stages.LIGHT_BLUE)
setp(6, 'Open questions remain unhighlighted. Automated tests and builds are not client acceptance. Jon’s latest confirmed passes have moved to Alex at his request; no new commit is implied. Only Alex’s returned green text marks an item Agreed. Financial checks remain with the client accountant.')
alex(39,40,'Jon has completed save/break and file-warning testing. Check Sheet imbalances remain soft findings: normal Save and Save As stay available, and recovery follows the chosen continue/pause option. Alex should repeat the saved-break and unsaved-discard sequences on a copy. This does not sign off the live header-refresh issue below.')
jon(43,'2.93 bounds the company-name editor to 700 logical pixels, while allowing it to shrink with the window. Check a long company name and a restored window at your normal display scale.')
stages.format_text(p[42],fill=stages.BLUE)
setp(44,'Check Sheet tab and public company-heading warning — still open')
setp(45,'Jon’s save/break and remembered file-warning sequences passed and are ready for Alex. An unsaved test finding must not alter the source workbook’s XML; the saved/recovery copy should carry only its own saved state.\n> Alex to test — Test response 2', stages.LIGHT_BLUE)
setp(46,'Unresolved: after a warm switch to Check Sheet, the displayed check updated quickly but the public company-heading warning did not. Jon also reports a value not clearing in Transactional DB; the exact cell and reproduction are still needed. Do not treat either result as resolved. The separate opt-in Check Sheet trial now uses a full calculation, not a worksheet-only calculation; formula/reference integrity remains a different check.')
jon(47,'Remaining separate test: Yes/No override clearing, Undo/Redo, and the Check Sheet trial’s live warning update. Record the input cell, Check Sheet row, Transactional DB cell and whether the trial ran. Turn temporary timings off after this trial is accepted; the reminder remains applicable.')
jon(50,'2.93 carries the configured 42-character minimum into the pivoted Opening Stock data path. Check that Stock Description is now wider and still resizable.')
jon(52,'The native grid test confirms the read-only separator immediately after Current Stock Numbers, before the pre-plan stock columns. Confirm it is visible on your display; no worksheet column is inserted.')
for i in (49,51): stages.format_text(p[i],fill=stages.BLUE)
setp(57,'Formatting on Click to access Economic Assumptions is indented.\nJon confirms the indenting is now satisfactory. Alex should check the link alignment on his display.\n> Alex to test — Test response 2', stages.LIGHT_BLUE)
heading_run = p[57].runs[0]
heading_text, detail_text = heading_run.text.split('\n', 1)
heading_run.text = heading_text + '\n'
detail_run = p[57].add_run(detail_text)
heading_run._r.addnext(detail_run._r)
stages.format_text(p[57], fill=stages.LIGHT_BLUE)
heading_run.font.color.rgb = RGBColor.from_string(stages.ORANGE)
for node in list(heading_run._r.get_or_add_rPr()):
    if node.tag in {qn('w:shd'), qn('w:highlight')}:
        node.getparent().remove(node)
alex(62,63,'Jon confirms the indenting is satisfactory. Alex should check the corresponding Specific Income link alignment.')
alex(72,73,'Jon confirms Management Costs formatting is satisfactory. Alex should check that the read-only category text remains readable on his display. This approval does not cover every conditional-rule or locking case elsewhere.')
stages.format_text(p[153],fill=stages.BLUE)
jon(154,'2.93 changes Facility Name from the Funders list to the workbook’s Facility list. The separate Funder editor still uses Funders. Compare both dropdowns with Excel.')
setp(200,'Current functional review')
setp(201,'Automated checks exercise disposable workbook copies. They are evidence for testing, not Jon’s approval or Alex’s agreement. Originals and repository master workbooks are not edited. Workbook fill-based locking, protected/formula cells and the existing change-history path remain in force.')
setp(202,'Priorities still open: Check Sheet public-header refresh, the reported Transactional DB remainder, unconfirmed layouts/nested scrolling and the questions under each issue. Shift+wheel grid zoom is feasible as a presentation feature, but is not enabled in this release; it needs a font/row-height pass that preserves workbook-owned formatting.')
setp(203,'Test the blue items below before client handoff. Older blue items remain outstanding unless separately confirmed. Keep the established functional-test numbering; the new tests continue at 61. Tests 57–59 for schedule focus, tinting and fixed figures remain relevant when checking the simplified scheduler.')
setp(204,'The earlier scaling-crash regression remains blue pending Jon’s functional result. No new claim of client acceptance is made by this revision.')

d.add_page_break()
d.add_heading('Jon functional tests for 2.93',1)
d.add_paragraph('New repairs remain blue. Record pass or fail beside each number, with the workbook and a screenshot where useful. Use a disposable copy for schedule tests.')
items = [
 ('61 Company name width','Open Global Assumptions on the large screen and in a restored window. The company-name input should be wider than a short date field but no longer fill almost the whole page. Edit a long name and confirm existing header/summary refresh.'),
 ('62 Opening Stock layout','Check the wider Stock Description column and the blank separator after Current Stock Numbers. Enter and clear a description: dependent cells must still unlock and relock according to workbook fill rules.'),
 ('63 Dropdown first click and keyboard','Click a dropdown arrow once: it should stay open. Test Shift+Down, Alt+Down and F4 on a cell and a repeating header editor. Read-only cells must remain protected. Then check normal arrow and Enter navigation.'),
 ('64 Funding wheel and context menu','Ordinary wheel should scroll vertically. Ctrl+wheel should move between loan columns. Right-click a loan cell and choose Add schedule…: the correct loan must be selected. Confirm the extra top header bar is gone; section schedule actions remain. Shift+wheel zoom is not part of this test.'),
 ('65 Facility Name choices','Compare Facility Name with Excel’s facility choices and Funder with its funder choices. They must not share the wrong list.'),
 ('66 Simple recurring dates','Set start and inclusive end dates. Test Monthly, Quarterly, Semi-annually and Annually. Annual 24 September should stay 24 September; 31 January should use February’s last valid day and return to 31 March. The preview must show one date column, with no working-day or holiday controls.'),
 ('67 Fixed schedule figures and preservation','Type 123456.789, then a negative decimal, in the fixed-figure input. Apply to a chosen loan and compatible sections. Dates should unlock the eligible amount cells, figures should appear only in that loan, and focus/tints should identify the new schedule. Check duplicate dates, one Undo/Redo, save and Excel/Summit reopen. No source-fill or lock overrides are permitted.'),
]
for heading,detail in items:
    h=d.add_heading(heading,2); stages.format_text(h,fill=stages.BLUE)
    q=d.add_paragraph(detail+'\n> Jon to test'); stages.format_text(q,fill=stages.BLUE)
    q.paragraph_format.keep_together=True
    q.paragraph_format.space_after=Pt(8)
d.add_heading('Open evidence to collect',2)
d.add_paragraph('Check Sheet and Transactional DB: supply the changed input address, expected Check Sheet row, stale Transactional DB address and the steps between them. A fast visible refresh alone does not prove the full calculation or the public warning state is current.')
output.parent.mkdir(parents=True,exist_ok=True)
d.save(output)
assert source.read_bytes()==before
with ZipFile(source) as a, ZipFile(output) as b:
    images=[n for n in a.namelist() if n.startswith('word/media/')]
    assert all(a.read(n)==b.read(n) for n in images), 'Original screenshots changed'
assert len(Document(output).inline_shapes)==len(Document(source).inline_shapes)
print('OUTPUT='+str(output.resolve()))
print('SOURCE_SHA256='+sha256(before).hexdigest())
print('OUTPUT_SHA256='+sha256(output.read_bytes()).hexdigest())
