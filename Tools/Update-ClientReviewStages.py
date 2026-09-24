"""Version the Word review with separate Jon/Alex test gates; preserve originals."""
import argparse
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

SOURCE_HASH = '9ba28ff78646fe3170f056d82b388c78c8d853db0ff1294fd122502f283c4fef'
BLUE = '9DC3E6'
LIGHT_BLUE = 'DDEBF7'
ORANGE = 'C45D18'
DARK = '222222'
GREEN = '196B24'
JON = '> Jon to test'
ALEX = '> Alex to test'
JON_DETAILS = {35, 38, 40, 41, 42, 45, 47, 52, 53, 58, 59, 61, 68, 71, 73, 77}
JON_ISSUES = {34, 37, 39, 44, 46, 57, 60, 67, 70, 72, 76}
ALEX_PAIRS = {
    7: 6, 9: 8, 11: 10, 13: 12, 15: 14, 17: 16, 19: 18,
    23: 22, 25: 24, 27: 26, 29: 28, 31: 30, 33: 32,
    50: 49, 55: 54, 65: 64, 75: 74, 83: 82, 96: 95,
    104: 103, 108: 107, 117: 116, 127: 126, 143: 142,
    153: 152, 162: 161, 164: 163, 175: 174, 185: 184,
    188: 187, 191: 190,
}


def format_text(p, colour=DARK, fill=None):
    for r in p.runs:
        if not r.text:
            continue
        r.font.color.rgb = RGBColor.from_string(colour)
        rpr = r._r.get_or_add_rPr()
        for node in list(rpr):
            if node.tag in {qn('w:shd'), qn('w:highlight')}:
                rpr.remove(node)
        if fill:
            shade = OxmlElement('w:shd')
            shade.set(qn('w:val'), 'clear')
            shade.set(qn('w:fill'), fill)
            rpr.append(shade)


def replace(p, text):
    drawings = [deepcopy(r._r) for r in p.runs if r._r.xpath('.//w:drawing')]
    for node in drawings:
        for t in node.xpath('.//w:t'):
            t.getparent().remove(t)
    p.clear()
    p.add_run(text)
    for node in drawings:
        p.add_run().add_break()
        p._p.append(node)


def after(p, text):
    el = OxmlElement('w:p')
    p._p.addnext(el)
    result = Paragraph(el, p._parent)
    result.style = 'Normal'
    result.paragraph_format.space_after = Pt(6)
    result.paragraph_format.keep_together = True
    result.paragraph_format.widow_control = True
    result.add_run(text)
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--client-original', required=True, type=Path)
    args = parser.parse_args()
    if args.output.exists():
        raise SystemExit('Output exists; choose a new revision.')
    before = args.source.read_bytes()
    assert sha256(before).hexdigest() == SOURCE_HASH, 'Review source changed'
    d = Document(args.source)
    ps = {i+1: p for i, p in enumerate(d.paragraphs)}
    assert len(ps) == 195 and 'version 03' in ps[1].text
    assert 'Blank button' in ps[20].text and 'Journals insertion' in ps[191].text
    client = Document(args.client_original)
    original_green = [p.text for p in client.paragraphs if any(
        str(r.font.color.rgb) == GREEN for r in p.runs if r.text)]
    assert ps[20].text in original_green, 'Existing green is not in original client file'
    # Remove the superseded yellow/ambiguous orange styling, not the source sign-off.
    for index, p in ps.items():
        if index != 20:
            format_text(p)
    replace(ps[1], 'Summit client issue review version 04')
    replace(ps[2], 'Ready to test with Summit 2.85\n23 September 2026\n\nJon tests the latest changes before handing them to Alex. The issue order, screenshots and open questions are retained. Word remains the working review so each issue can carry its evidence and discussion.')
    replace(ps[3], 'Test stages and colour key')
    ps[3].runs[0].bold = True
    key_jon = after(ps[3], 'Blue: Jon to test — uncommitted 2.84 changes and the 2.85 scaling-crash repair require Jon\'s functional test.')
    format_text(key_jon, fill=BLUE)
    key_alex = after(key_jon, 'Orange issue heading + light-blue detail: Alex to test — committed fixes ready for client testing, not signed off.')
    format_text(key_alex, fill=LIGHT_BLUE)
    key_alex.runs[0].font.color.rgb = RGBColor.from_string(ORANGE)
    key_green = after(key_alex, 'Green font: Agreed — only when Alex returns his Word document with the item green. Existing client green is retained.')
    format_text(key_green, GREEN)
    after(key_green, 'Open questions and unresolved parts remain unhighlighted. A scope decision, successful automated test, build or commit is not client acceptance. Move the latest fixes to Alex only after Jon\'s functional pass and commit. Financial checks remain with the client accountant.')
    replace(ps[4], 'Client issues from Summit 2.51')
    crash = after(ps[5], 'Interface scaling — unhandled collection-modified exception while forms are opening or closing (reported in 2.84).')
    crash.paragraph_format.keep_with_next = True
    format_text(crash, fill=BLUE)
    crash_test = after(crash, 'Repair in 2.85: scaling retries a changing open-window list, defers safely if it remains busy, and applies each form\'s outstanding scale only once. Main-thread coordination avoids repeated cross-thread idle callbacks. Please repeat the action that triggered the crash, open/close interfaces and Options, and check 100% → 125% → 100% scaling.\n' + JON)
    format_text(crash_test, fill=BLUE)
    # Keep source green and withdraw the generated question against that signed-off item.
    replace(ps[21], 'Agreed — green retained from the original client document.')
    format_text(ps[21], GREEN)
    replace(ps[62], 'Scope decision (not test acceptance): remove expansion from this interface definition; do not add a new workbook expansion rule.')
    replace(ps[78], 'Scope decision (not test acceptance): pin the category columns as well as Description.')
    replace(ps[71], 'Partial repair to test: the two category fields now refresh their cached source colours and rule state, rather than retaining stale colours.\n' + JON + ' — this limited repair only.')
    after(ps[71], 'Still open: full conditional-format parity with Excel. Question: after clearing Description and committing the edit, which colour or editability still differs from the corresponding Excel cell?')
    for index in JON_DETAILS:
        p = ps[index]
        text = p.text.replace(ALEX, JON).strip()
        if JON not in text:
            text += '\n' + JON
        replace(p, text)
        format_text(p, fill=BLUE)
    for index in JON_ISSUES:
        format_text(ps[index], fill=BLUE)
    for detail, issue in ALEX_PAIRS.items():
        assert ALEX in ps[detail].text, f'Unexpected client status: {detail}'
        format_text(ps[issue], ORANGE)
        format_text(ps[detail], fill=LIGHT_BLUE)
    # This source paragraph contains both open layout work and a committed input repair.
    replace(ps[84], 'Survey Input — headers need wrapping and columns narrowing. Scroll within a scroll.\n\nCannot seem to paste into input cells (not allowing decimal places?). This needs to support pasting from an Excel table.')
    format_text(ps[84])
    replace(ps[85], 'Committed partial repair: Survey cost inputs retain decimal entry and paste values. Please retest a rectangular paste from Excel and identify any failing destination cells.\n' + ALEX + ' — decimal entry and paste only.')
    partial_heading = ps[85].insert_paragraph_before('Survey Input — decimal entry and paste')
    partial_heading.paragraph_format.keep_with_next = True
    format_text(partial_heading, ORANGE)
    format_text(ps[85], fill=LIGHT_BLUE)
    after(ps[85], 'Still open: Survey Input header wrapping, column widths and nested scrolling. The input repair does not resolve these layout items.')
    replace(ps[192], 'Jon functional test before client handoff')
    replace(ps[193], 'Debug and Release each passed 35 native assertions on private workbook copies. They cover Yes/No entry and Undo/Redo, session-only Check Sheet findings, save persistence, recovery policy, automatic recheck and idle scheduling, pinned Management Costs descriptors, source-colour refresh and the wider company-name editor. These are automated checks, not Jon\'s functional approval or Alex\'s sign-off. Original master files are unchanged.')
    replace(ps[195], 'Jon: test the blue items in 2.85, which includes the pending 2.84 changes. Make an imbalance without saving, check it, close and discard, then reopen. Repeat with a deliberate Save. Confirm the message follows the saved file; test Yes/No clearing and Undo/Redo, recovery continuation and the trial settings. Check the blue layout and readability changes at normal display scale. Record pass/fail and screenshots beside each issue. Formula errors must remain Integrity findings. After the functional pass and commit, change those items to the Alex stage; only Alex\'s returned green marks agreement.')
    after(ps[195], 'Scaling regression: 16 focused checks pass in both Debug and Release, including a forced collection-modified exception, deferred and re-entrant scaling, disposal, and 200 real form open/close cycles on a second UI thread. This remains blue pending Jon\'s functional test.')
    for p in d.paragraphs:
        p.paragraph_format.widow_control = True
        p.paragraph_format.keep_together = True
        if p.style.name.startswith(('Title', 'Heading', 'Subtitle')):
            p.paragraph_format.keep_with_next = True
    for issue in JON_ISSUES | set(ALEX_PAIRS.values()):
        ps[issue].paragraph_format.keep_with_next = True
    # Machine checks prevent accidental promotion or residual yellow test markers.
    assert all(JON in ps[i].text and ALEX not in ps[i].text for i in JON_DETAILS)
    assert all(ALEX in ps[i].text and JON not in ps[i].text for i in ALEX_PAIRS)
    assert 'yellow' not in '\n'.join(p.text for p in d.paragraphs).lower()
    assert not d._element.xpath('.//w:shd[@w:fill="FFFF99"]')
    args.output.parent.mkdir(parents=True, exist_ok=True)
    d.save(args.output)
    assert sha256(args.source.read_bytes()).hexdigest() == SOURCE_HASH
    with ZipFile(args.source) as src, ZipFile(args.output) as out:
        media = [n for n in src.namelist() if n.startswith('word/media/')]
        assert all(src.read(n) == out.read(n) for n in media)
    print(f'Created {args.output}; {len(JON_DETAILS)} Jon test notes; {len(ALEX_PAIRS)+1} Alex test notes; {len(media)} original images unchanged; source unchanged.')


if __name__ == '__main__':
    main()
